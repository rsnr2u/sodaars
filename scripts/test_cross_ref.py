import openpyxl

wb = openpyxl.load_workbook(r"c:\xampp\htdocs\DIGITALKSIND\sodaars\SODAARS_Project_Links_Audit.xlsx")
ws1 = wb['All Project Links Audit']
ws2 = wb['Defined Routes Inventory']

extracted_urls = set()
for row in ws1.iter_rows(min_row=2, values_only=True):
    if row and row[6]:
        extracted_urls.add(str(row[6]).split('?')[0].split('#')[0].rstrip('/'))

defined_routes = set()
for row in ws2.iter_rows(min_row=2, values_only=True):
    if row and row[1]:
        defined_routes.add(str(row[1]).rstrip('/'))

print(f"Total defined routes: {len(defined_routes)}")
print(f"Total distinct link targets: {len(extracted_urls)}")

orphans = defined_routes - extracted_urls
print(f"Routes defined but NOT directly linked in static links: {len(orphans)}")
for o in sorted(orphans):
    print("  Orphan route:", o)
